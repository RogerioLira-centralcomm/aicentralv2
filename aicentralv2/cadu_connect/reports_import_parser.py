"""Bounded, platform-agnostic campaign export parsing. No database or AI calls."""

import csv
import io
import re
import unicodedata
import zipfile
from datetime import date, datetime
from decimal import Decimal, InvalidOperation

from openpyxl import load_workbook


MAX_FILE_BYTES = 5 * 1024 * 1024
MAX_UNCOMPRESSED_XLSX = 30 * 1024 * 1024
MAX_ROWS = 2000
MAX_COLUMNS = 80

ALIASES = {
    'platform': {'platform', 'plataforma', 'ad platform', 'media platform'},
    'account_id': {'account id', 'account_id', 'customer id', 'ad account id', 'id da conta', 'id conta', 'id da conta de anuncios'},
    'account_name': {'account', 'account name', 'ad account name', 'nome da conta', 'conta'},
    'campaign_id': {'campaign id', 'campaign_id', 'id da campanha', 'id campanha'},
    'campaign_name': {'campaign', 'campaign name', 'nome da campanha', 'campanha'},
    'date': {'date', 'day', 'data', 'dia', 'reporting date'},
    'currency': {'currency', 'currency code', 'moeda', 'codigo da moeda'},
    'impressions': {'impressions', 'impressoes', 'impression'},
    'clicks': {'clicks', 'cliques', 'click'},
    'cost': {'cost', 'spend', 'amount spent', 'custo', 'gasto', 'valor gasto'},
    'conversions': {'conversions', 'conversoes', 'conversion'},
    'conversion_value': {'conversion value', 'conversions value', 'valor de conversao', 'valor das conversoes'},
}
FIELD_BY_HEADER = {alias: field for field, aliases in ALIASES.items() for alias in aliases}
METRICS = ('impressions', 'clicks', 'cost', 'conversions', 'conversion_value')
PLATFORMS = {
    'google ads': 'google_ads', 'adwords': 'google_ads', 'meta ads': 'meta_ads',
    'facebook ads': 'meta_ads', 'instagram ads': 'meta_ads',
    'microsoft ads': 'microsoft_ads', 'bing ads': 'microsoft_ads',
    'tiktok ads': 'tiktok_ads', 'linkedin ads': 'linkedin_ads',
    'pinterest ads': 'pinterest_ads',
}


def normalized_header(value):
    value = unicodedata.normalize('NFKD', str(value or '').casefold())
    value = ''.join(char for char in value if not unicodedata.combining(char))
    return ' '.join(re.sub(r'[^a-z0-9]+', ' ', value).split())


def normalized_platform(value):
    text = normalized_header(value)
    if not text:
        return ''
    slug = PLATFORMS.get(text) or text.replace(' ', '_')
    return slug if re.fullmatch(r'[a-z][a-z0-9_]{0,31}', slug) else ''


def _headers(values):
    if not values or len(values) > MAX_COLUMNS:
        raise ValueError('O arquivo precisa ter de 1 a 80 colunas.')
    names = [str(value or '').strip()[:160] for value in values]
    if not any(names):
        raise ValueError('Cabeçalho vazio.')
    present = [normalized_header(name) for name in names if name]
    if len(present) != len(set(present)):
        raise ValueError('Há cabeçalhos repetidos no arquivo.')
    mapped = [FIELD_BY_HEADER.get(normalized_header(name)) for name in names]
    known = [field for field in mapped if field]
    if len(known) != len(set(known)):
        raise ValueError('Há colunas equivalentes repetidas; escolha uma coluna por campo.')
    return names


def _record(headers, values, sheet, row_number):
    if len(values) > len(headers) and any(value not in ('', None) for value in values[len(headers):]):
        raise ValueError(f'{sheet}, linha {row_number}: há mais valores do que colunas.')
    return {'sheet': sheet, 'row': row_number,
            'raw': {header: _cell(values[index]) if index < len(values) else ''
                    for index, header in enumerate(headers) if header}}


def _cell(value):
    if value is None:
        return ''
    if isinstance(value, (date, datetime)):
        return value.date().isoformat() if isinstance(value, datetime) else value.isoformat()
    if isinstance(value, float) and value.is_integer() and abs(value) <= 9_000_000_000_000_000:
        return str(int(value))
    return str(value).strip()[:1000]


def read_export(raw, filename):
    if not raw or len(raw) > MAX_FILE_BYTES:
        raise ValueError('O arquivo deve ter até 5 MiB e não pode estar vazio.')
    name = str(filename or '').lower()
    if name.endswith('.csv'):
        try:
            content = raw.decode('utf-8-sig')
        except UnicodeDecodeError:
            content = raw.decode('cp1252')
        sample = content[:4096]
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=',;\t')
        except csv.Error:
            dialect = csv.excel
        reader = csv.reader(io.StringIO(content), dialect)
        try:
            headers = _headers(next(reader))
        except StopIteration as exc:
            raise ValueError('CSV vazio.') from exc
        records = []
        for number, values in enumerate(reader, start=2):
            if not any(str(value or '').strip() for value in values):
                continue
            if len(records) >= MAX_ROWS:
                raise ValueError('O arquivo excede 2.000 linhas de dados.')
            records.append(_record(headers, values, 'CSV', number))
        if not records:
            raise ValueError('O CSV não contém linhas de dados.')
        return 'csv', records
    if name.endswith('.xlsx'):
        if not zipfile.is_zipfile(io.BytesIO(raw)):
            raise ValueError('XLSX inválido.')
        with zipfile.ZipFile(io.BytesIO(raw)) as archive:
            if sum(item.file_size for item in archive.infolist()) > MAX_UNCOMPRESSED_XLSX:
                raise ValueError('O XLSX descompactado excede 30 MiB.')
        workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=True, keep_links=False)
        records = []
        try:
            for sheet in workbook.worksheets:
                headers = None
                for number, values in enumerate(sheet.iter_rows(values_only=True), start=1):
                    if not any(value not in ('', None) for value in values):
                        continue
                    if headers is None:
                        headers = _headers(values)
                        continue
                    if len(records) >= MAX_ROWS:
                        raise ValueError('O arquivo excede 2.000 linhas de dados.')
                    records.append(_record(headers, values, sheet.title[:100], number))
        finally:
            workbook.close()
        if not records:
            raise ValueError('O XLSX não contém linhas de dados.')
        return 'xlsx', records
    raise ValueError('Envie CSV ou XLSX para leitura tabular.')


def _decimal(value):
    raw = str(value or '').strip().replace('\u00a0', ' ').replace('R$', '').replace('$', '').replace(' ', '')
    if not raw:
        return None
    if raw.startswith('-') or '%' in raw:
        raise ValueError('valor negativo ou percentual')
    if ',' in raw and '.' in raw:
        decimal_mark = ',' if raw.rfind(',') > raw.rfind('.') else '.'
        grouping_mark = '.' if decimal_mark == ',' else ','
        if not re.fullmatch(r'\d{1,3}(?:' + re.escape(grouping_mark) + r'\d{3})+' +
                            re.escape(decimal_mark) + r'\d{1,6}', raw):
            raise ValueError('separadores numéricos ambíguos')
        raw = raw.replace(grouping_mark, '').replace(decimal_mark, '.')
    elif ',' in raw or '.' in raw:
        mark = ',' if ',' in raw else '.'
        if not re.fullmatch(r'\d+' + re.escape(mark) + r'\d{1,6}', raw):
            raise ValueError('número inválido')
        if len(raw.rsplit(mark, 1)[1]) == 3:
            raise ValueError('separador de milhar ou decimal ambíguo')
        raw = raw.replace(mark, '.')
    elif not raw.isdigit():
        raise ValueError('número inválido')
    try:
        result = Decimal(raw)
    except InvalidOperation as exc:
        raise ValueError('número inválido') from exc
    if not result.is_finite() or result >= Decimal('1000000000000000000'):
        raise ValueError('número fora do limite')
    return result


def _date(value, order):
    raw = str(value or '').strip()
    if not raw:
        return None
    try:
        return date.fromisoformat(raw).isoformat()
    except ValueError:
        pass
    parts = re.fullmatch(r'(\d{1,2})/(\d{1,2})/(\d{4})', raw)
    if not parts:
        raise ValueError('data não reconhecida')
    first, second, year = map(int, parts.groups())
    if order == 'auto' and first <= 12 and second <= 12:
        raise ValueError('dia e mês ambíguos')
    inferred_mdy = order == 'auto' and second > 12 and first <= 12
    day, month = (second, first) if order == 'mdy' or inferred_mdy else (first, second)
    try:
        return date(year, month, day).isoformat()
    except ValueError as exc:
        raise ValueError('data inválida') from exc


def parse_record(record, *, platform_hint='', currency_hint='', date_order='auto'):
    values = {}
    for header, value in record['raw'].items():
        field = FIELD_BY_HEADER.get(normalized_header(header))
        if field:
            values[field] = value
    issues = []
    platform = normalized_platform(values.get('platform') or platform_hint)
    account_id = str(values.get('account_id') or '').strip()
    account_name = str(values.get('account_name') or '').strip()
    campaign_id = str(values.get('campaign_id') or '').strip()
    campaign_name = str(values.get('campaign_name') or '').strip()
    if platform == 'google_ads' and re.fullmatch(r'[\d-]+', account_id):
        account_id = account_id.replace('-', '')
    for field, value, limit in (('ID da conta', account_id, 160), ('nome da conta', account_name, 240),
                                ('ID da campanha', campaign_id, 160), ('nome da campanha', campaign_name, 240)):
        if not value or len(value) > limit:
            issues.append(f'{field} ausente ou longo demais')
    if not platform:
        issues.append('plataforma não identificada')
    try:
        metric_date = _date(values.get('date'), date_order)
    except ValueError as exc:
        metric_date = None
        issues.append(str(exc))
    if not metric_date:
        issues.append('data ausente')
    currency = str(values.get('currency') or currency_hint or '').strip().upper()
    if currency and not re.fullmatch(r'[A-Z]{3}', currency):
        issues.append('moeda inválida')
        currency = ''
    metrics = {}
    for field in METRICS:
        if field not in values or not str(values[field]).strip():
            continue
        try:
            number = _decimal(values[field])
            if field in ('impressions', 'clicks') and number != number.to_integral_value():
                raise ValueError('contagem fracionária')
            metrics[field] = str(number)
        except ValueError as exc:
            issues.append(f'{field}: {exc}')
    if not metrics:
        issues.append('nenhuma métrica conhecida')
    if any(field in metrics for field in ('cost', 'conversion_value')) and not currency:
        issues.append('moeda necessária para valores monetários')
    return {
        'platform': platform, 'external_account_id': account_id, 'account_name': account_name,
        'external_campaign_id': campaign_id, 'campaign_name': campaign_name,
        'metric_date': metric_date, 'currency': currency, 'metrics': metrics,
        'issues': list(dict.fromkeys(issues)),
    }
