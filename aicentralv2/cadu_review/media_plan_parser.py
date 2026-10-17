"""Lê um plano de mídia de qualquer agência (XLSX, CSV ou PDF) em linhas normalizadas.

Planos não têm padrão: o cabeçalho pode estar na linha 8, a verba pode vir bruta e
líquida, o modelo de compra pode estar implícito no nome da coluna ("CPM"). O leitor
procura o cabeçalho pelos sinônimos abaixo, guarda o par rótulo/valor acima dele
(objetivo, público, verba…) como contexto e separa as linhas de total.
"""
from __future__ import annotations

import csv
import io
import re
import unicodedata
import zipfile
from datetime import date, datetime

MAX_FILE_BYTES = 5 * 1024 * 1024
MAX_UNCOMPRESSED_XLSX = 30 * 1024 * 1024
MAX_ROWS = 300
MAX_PDF_PAGES = 20
HEADER_SCAN_ROWS = 40

FIELD_ALIASES = {
    'canal': ('veiculo', 'canal', 'midia', 'meio', 'plataforma', 'publisher', 'portal', 'site', 'rede',
              'emissora', 'player', 'parceiro', 'fornecedor', 'veiculo midia'),
    'formato': ('formato', 'peca', 'tipo de midia', 'entrega', 'posicionamento', 'produto', 'formato peca',
                'tipo de formato', 'criativo'),
    'modelo': ('modelo de compra', 'modelo', 'compra', 'negociacao', 'tipo de compra', 'forma de compra',
               'unidade', 'unidade de compra', 'metrica de compra', 'tipo de negociacao'),
    'custo_unitario': ('custo unitario', 'valor unitario', 'preco unitario', 'unitario', 'valor unit',
                       'custo unit', 'cpm', 'cpc', 'cpv', 'cpa', 'cpl', 'cpe', 'cpm negociado', 'custo',
                       'valor negociado unitario', 'cpm liquido', 'cpc liquido'),
    'quantidade': ('quantidade', 'qtd', 'qtde', 'volume', 'impressoes', 'impressoes previstas', 'cliques previstos',
                   'views', 'visualizacoes', 'entregas', 'insercoes', 'volume previsto', 'entrega prevista',
                   'impressoes contratadas', 'cliques', 'leads previstos'),
    'investimento': ('investimento', 'valor total', 'custo total', 'valor liquido', 'liquido', 'budget', 'verba',
                     'valor negociado', 'investimento liquido', 'investimento total', 'total liquido', 'valor',
                     'total investimento', 'investimento r'),
    'investimento_bruto': ('valor bruto', 'bruto', 'investimento bruto', 'total bruto', 'valor tabela'),
    'percentual': ('percentual', 'share', 'participacao', 'perc', 'percentual verba', 'do total', 'peso'),
    'inicio': ('inicio', 'data inicio', 'data de inicio', 'start', 'data inicial', 'de'),
    'fim': ('fim', 'termino', 'data fim', 'data de termino', 'end', 'data final', 'ate'),
    'periodo': ('periodo', 'flight', 'voo', 'veiculacao', 'periodo de veiculacao'),
    'kpi': ('kpi', 'kpi principal', 'meta', 'objetivo da linha', 'objetivo'),
    'publico': ('publico', 'target', 'segmentacao', 'audiencia', 'publico alvo'),
    'praca': ('praca', 'regiao', 'geografia', 'cidade', 'uf', 'mercado'),
}
CONTEXT_ALIASES = {
    'cliente': ('cliente', 'anunciante', 'marca'),
    'campanha': ('campanha', 'projeto', 'acao'),
    'objetivo': ('objetivo', 'objetivo da campanha', 'objetivo de comunicacao'),
    'publico': ('publico', 'publico alvo', 'target', 'audiencia'),
    'kpi': ('kpi', 'kpi principal', 'meta', 'metas'),
    'periodo': ('periodo', 'periodo da campanha', 'veiculacao', 'flight'),
    'verba': ('verba', 'verba total', 'investimento total', 'budget', 'investimento'),
    'praca': ('praca', 'pracas', 'regiao', 'geografia'),
}
TOTAL_LABELS = {'total', 'subtotal', 'total geral', 'total liquido', 'total bruto', 'total da campanha',
                'total investimento', 'grand total', 'soma'}
MODEL_WORDS = {
    'cpm': ('cpm', 'mil impressoes', 'impressao', 'impressoes'),
    'cpc': ('cpc', 'clique', 'cliques'),
    'cpv': ('cpv', 'cpcv', 'view', 'views', 'visualizacao', 'video completo'),
    'cpa': ('cpa', 'cpl', 'lead', 'conversao', 'aquisicao'),
    'cpe': ('cpe', 'engajamento'),
    'fixo': ('fixo', 'flat', 'pacote', 'diaria', 'mensal', 'insercao', 'insercoes', 'patrocinio', 'spot',
             'semanal', 'cota', 'projeto especial'),
}
_ALIAS_FIELD = {alias: field for field, aliases in FIELD_ALIASES.items() for alias in aliases}
_CONTEXT_FIELD = {alias: field for field, aliases in CONTEXT_ALIASES.items() for alias in aliases}
# Curtos ou genéricos demais para casar por prefixo ("de", "até", "valor"…).
_EXACT_ONLY = {'ate', 'de', 'qtd', 'valor', 'meta', 'custo', 'site', 'rede', 'meio', 'uf', 'perc', 'peso', 'compra'}


class PlanParseError(ValueError):
    """Mensagem segura para mostrar à pessoa que enviou o arquivo."""


def normalized(value) -> str:
    text = unicodedata.normalize('NFKD', str(value or '').casefold())
    text = ''.join(char for char in text if not unicodedata.combining(char))
    return ' '.join(re.sub(r'[^a-z0-9]+', ' ', text).split())


def number(value):
    """Número de planilha brasileira ou americana: 'R$ 1.234,56', '1,234.56', '15%', 12000."""
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip()
    if not text or text in {'-', '—', '–'}:
        return None
    negative = text.startswith('(') and text.endswith(')') or text.startswith('-')
    text = re.sub(r'[^0-9,.]', '', text)
    if not text or not re.search(r'\d', text):
        return None
    if ',' in text and '.' in text:
        if text.rfind(',') > text.rfind('.'):
            text = text.replace('.', '').replace(',', '.')
        else:
            text = text.replace(',', '')
    elif ',' in text:
        text = text.replace('.', '').replace(',', '.') if text.count(',') == 1 else text.replace(',', '')
    elif text.count('.') > 1 or re.fullmatch(r'\d{1,3}(\.\d{3})+', text):
        text = text.replace('.', '')
    try:
        result = float(text)
    except ValueError:
        return None
    return -result if negative else result


def _date(value, year_hint=None):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, (int, float)) and 30000 < value < 60000:  # serial do Excel
        return date.fromordinal(date(1899, 12, 30).toordinal() + int(value))
    text = str(value or '').strip()
    found = re.fullmatch(r'(\d{4})-(\d{1,2})-(\d{1,2})(?:[ T].*)?', text)
    if found:
        y, m, d = (int(part) for part in found.groups())
    else:
        found = re.fullmatch(r'(\d{1,2})[/.-](\d{1,2})(?:[/.-](\d{2,4}))?', text)
        if not found:
            return None
        d, m = int(found.group(1)), int(found.group(2))
        y = int(found.group(3)) if found.group(3) else year_hint
        if y is None:
            return None
        if y < 100:
            y += 2000
    try:
        return date(y, m, d)
    except ValueError:
        return None


def period(value):
    """'01/11 a 30/11/2026', '01/11/2026 - 15/12/2026' → (início, fim)."""
    text = str(value or '')
    parts = re.findall(r'\d{4}-\d{1,2}-\d{1,2}|\d{1,2}[/.]\d{1,2}(?:[/.]\d{2,4})?', text)
    if len(parts) < 2:
        return None, None
    end = _date(parts[-1])
    start = _date(parts[0], year_hint=end.year if end else None)
    return start, end


def buying_model(*values):
    for value in values:
        words = normalized(value)
        if not words:
            continue
        for model, keys in MODEL_WORDS.items():
            if any(re.search(rf'\b{re.escape(key)}\b', words) for key in keys):
                return model
    return ''


def _cell(value):
    if value is None:
        return ''
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer() and abs(value) < 1e15:
        return str(int(value))
    return str(value).strip()[:500]


def _header_fields(values):
    """Mapeia cada coluna para um campo; o primeiro sinônimo vence quando dois disputam o mesmo campo."""
    fields, used = {}, set()
    for index, value in enumerate(values):
        name = normalized(value)
        field = _ALIAS_FIELD.get(name)
        if not field:
            field = next((candidate for alias, candidate in _ALIAS_FIELD.items()
                          if len(alias) >= 3 and alias not in _EXACT_ONLY
                          and (name.startswith(alias + ' ') or name.endswith(' ' + alias))), None)
        if field and field not in used:
            fields[index] = field
            used.add(field)
    return fields


def _header_score(fields):
    found = set(fields.values())
    if 'canal' not in found or not found.intersection({'investimento', 'investimento_bruto', 'custo_unitario', 'quantidade'}):
        return 0
    return len(found)


def _context_pairs(rows, limit):
    context = {}
    for values in rows[:limit]:
        cells = [_cell(value) for value in values]
        for index, cell in enumerate(cells):
            if not cell:
                continue
            label, _, inline = cell.partition(':')
            field = _CONTEXT_FIELD.get(normalized(label))
            if not field or field in context:
                continue
            value = inline.strip() or next((item for item in cells[index + 1:] if item), '')
            if value:
                context[field] = value[:500]
    return context


def _table(rows, source):
    """Escolhe o cabeçalho com mais campos reconhecidos nas primeiras linhas."""
    best = None
    for index, values in enumerate(rows[:HEADER_SCAN_ROWS]):
        fields = _header_fields(values)
        score = _header_score(fields)
        if score >= 2 and (best is None or score > best[0]):
            best = (score, index, fields)
    if not best:
        return None
    _, header_index, fields = best
    header = [_cell(value) for value in rows[header_index]]
    context = _context_pairs(rows, header_index)
    lines, totals = [], []
    for offset, values in enumerate(rows[header_index + 1:], start=header_index + 2):
        record = {field: values[column] if column < len(values) else None for column, field in fields.items()}
        if not any(_cell(value) for value in record.values()):
            continue
        label = normalized(record.get('canal'))
        first = normalized(next((value for value in values if _cell(value)), ''))
        if label in TOTAL_LABELS or first in TOTAL_LABELS or first.startswith('total '):
            totals.append({'line': offset, 'investimento': number(record.get('investimento')),
                           'investimento_bruto': number(record.get('investimento_bruto'))})
            continue
        if not label and number(record.get('investimento')) is not None and not _cell(record.get('formato')):
            totals.append({'line': offset, 'investimento': number(record.get('investimento')), 'subtotal': True})
            continue
        if len(lines) >= MAX_ROWS:
            raise PlanParseError(f'O plano passa de {MAX_ROWS} linhas. Envie só a aba do plano.')
        lines.append(_line(record, offset, header, fields, source))
    return {'context': context, 'lines': lines, 'totals': totals, 'header': [h for h in header if h], 'source': source}


def _line(record, line_number, header, fields, source):
    unit_header = next((header[column] for column, field in fields.items() if field == 'custo_unitario'), '')
    start, end = _date(record.get('inicio')), _date(record.get('fim'))
    if (not start or not end) and record.get('periodo'):
        start, end = period(record.get('periodo'))
    if start and not end and record.get('fim'):
        end = _date(record.get('fim'), year_hint=start.year)
    share = number(record.get('percentual'))
    if share is not None and 0 < share <= 1 and '%' not in str(record.get('percentual') or ''):
        share *= 100
    return {
        'linha': line_number, 'aba': source,
        'canal': _cell(record.get('canal')), 'formato': _cell(record.get('formato')),
        'modelo': buying_model(record.get('modelo'), unit_header),
        'modelo_original': _cell(record.get('modelo')),
        'custo_unitario': number(record.get('custo_unitario')),
        'quantidade': number(record.get('quantidade')),
        'investimento': number(record.get('investimento')),
        'investimento_bruto': number(record.get('investimento_bruto')),
        'percentual': share,
        'inicio': start.isoformat() if start else '', 'fim': end.isoformat() if end else '',
        'kpi': _cell(record.get('kpi')), 'publico': _cell(record.get('publico')), 'praca': _cell(record.get('praca')),
    }


def _csv_rows(raw):
    try:
        content = raw.decode('utf-8-sig')
    except UnicodeDecodeError:
        content = raw.decode('cp1252')
    # O Sniffer escolhe a vírgula em planilhas brasileiras ("12.000,00"); ';' e tab presentes
    # na maioria das linhas são o separador.
    sample = [line for line in content.splitlines()[:40] if line.strip()]
    delimiter = next((candidate for candidate in (';', '\t')
                      if sum(candidate in line for line in sample) * 2 >= len(sample)), ',')
    return [('CSV', list(csv.reader(io.StringIO(content), delimiter=delimiter)))]


def _xlsx_rows(raw):
    from openpyxl import load_workbook
    if not zipfile.is_zipfile(io.BytesIO(raw)):
        raise PlanParseError('XLSX inválido.')
    with zipfile.ZipFile(io.BytesIO(raw)) as archive:
        if sum(item.file_size for item in archive.infolist()) > MAX_UNCOMPRESSED_XLSX:
            raise PlanParseError('O XLSX descompactado passa de 30 MiB.')
    workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=True, keep_links=False)
    try:
        return [(sheet.title[:80], [list(row) for _, row in zip(range(MAX_ROWS + HEADER_SCAN_ROWS + 50),
                                                                  sheet.iter_rows(values_only=True))])
                for sheet in workbook.worksheets]
    finally:
        workbook.close()


def _pdf_rows(raw):
    import fitz
    sheets, text = [], []
    with fitz.open(stream=raw, filetype='pdf') as document:
        if document.page_count > MAX_PDF_PAGES:
            raise PlanParseError(f'O PDF passa de {MAX_PDF_PAGES} páginas. Envie só as páginas do plano.')
        for page in document:
            text.append(page.get_text('text'))
            for number_, table in enumerate(page.find_tables().tables, start=1):
                rows = [[cell if cell is not None else '' for cell in row] for row in table.extract()]
                if rows:
                    sheets.append((f'Página {page.number + 1}, tabela {number_}', rows))
    preface = [[line] for line in '\n'.join(text).splitlines()[:80]]
    return sheets, preface


def parse(raw: bytes, filename: str) -> dict:
    """Devolve {'lines', 'context', 'totals', 'source', 'file_type'} ou levanta PlanParseError."""
    if not raw:
        raise PlanParseError('O arquivo está vazio.')
    if len(raw) > MAX_FILE_BYTES:
        raise PlanParseError('O arquivo deve ter até 5 MiB.')
    name = str(filename or '').lower()
    preface = []
    if name.endswith('.csv'):
        file_type, sheets = 'csv', _csv_rows(raw)
    elif name.endswith('.xlsx'):
        file_type, sheets = 'xlsx', _xlsx_rows(raw)
    elif name.endswith('.pdf'):
        file_type = 'pdf'
        sheets, preface = _pdf_rows(raw)
    else:
        raise PlanParseError('Envie o plano em XLSX, CSV ou PDF.')
    tables = [table for source, rows in sheets if (table := _table(rows, source))]
    tables = [table for table in tables if table['lines']]
    if not tables:
        raise PlanParseError('Não encontrei a tabela do plano. Confira se há colunas como Veículo/Canal e Investimento.')
    chosen = max(tables, key=lambda table: len(table['lines']))
    if preface:
        chosen['context'] = {**_context_pairs(preface, len(preface)), **chosen['context']}
    chosen['file_type'] = file_type
    return chosen
