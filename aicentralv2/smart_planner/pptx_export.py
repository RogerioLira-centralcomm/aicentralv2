"""Exporta a proposta em slides como PowerPoint editável, no modelo de apresentação do plano.

Usa os mesmos slides da página pública (``public_slides.build_slides``): o que
o cliente vê no link é o que sai no arquivo, com texto editável, gráficos em
formas simples e as notas do apresentador.
"""

from __future__ import annotations

import base64
import io
import os
from typing import Any

from .helpers import as_dict, as_list, text

W_IN, H_IN = 13.333, 7.5
PAD = 0.75
# Fonte presente em qualquer PowerPoint, Keynote ou Google Slides; o arquivo continua editável.
FONT = "Arial"


def _rgb(color: str):
    from pptx.dml.color import RGBColor

    value = text(color).lstrip("#") or "10252D"
    if len(value) == 3:
        value = "".join(ch * 2 for ch in value)
    return RGBColor.from_string(value[:6].upper())


def _image_stream(url: str, static_folder: str) -> io.BytesIO | None:
    """Só imagens locais ou embutidas: o export não baixa nada da internet."""
    value = text(url)
    if value.startswith("data:image/"):
        try:
            return io.BytesIO(base64.b64decode(value.partition(",")[2], validate=True))
        except Exception:
            return None
    if value.startswith("/static/") and static_folder:
        root = os.path.abspath(static_folder)
        path = os.path.abspath(os.path.join(root, value[len("/static/"):].split("?", 1)[0]))
        if path.startswith(root) and os.path.isfile(path) and not path.lower().endswith(".svg"):
            with open(path, "rb") as handle:
                return io.BytesIO(handle.read())
    return None


class _Deck:
    def __init__(self, template: dict, static_folder: str):
        from pptx import Presentation
        from pptx.util import Inches

        self.Inches = Inches
        self.prs = Presentation()
        self.prs.slide_width = Inches(W_IN)
        self.prs.slide_height = Inches(H_IN)
        self.blank = self.prs.slide_layouts[6]
        tokens = as_dict(template.get("tokens"))
        self.ink = tokens.get("ink") or "#10252D"
        self.brand = tokens.get("brand") or "#0B6B52"
        self.accent = tokens.get("accent") or "#F0B000"
        self.paper = tokens.get("paper") or "#FFFEFA"
        self.soft = tokens.get("soft") or "#EEF5F1"
        self.muted = tokens.get("muted") or "#5F7076"
        self.chart = [c for c in as_list(template.get("chart")) if text(c)] or [self.brand, self.accent]
        self.static_folder = static_folder

    # -- primitivas

    def slide(self, background: str | None = None):
        slide = self.prs.slides.add_slide(self.blank)
        fill = slide.background.fill
        fill.solid()
        fill.fore_color.rgb = _rgb(background or self.paper)
        return slide

    def box(self, slide, x, y, w, h, color: str, line: str | None = None):
        from pptx.enum.shapes import MSO_SHAPE

        shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, self.Inches(x), self.Inches(y), self.Inches(w), self.Inches(h))
        shape.fill.solid()
        shape.fill.fore_color.rgb = _rgb(color)
        if line:
            shape.line.color.rgb = _rgb(line)
        else:
            shape.line.fill.background()
        shape.shadow.inherit = False
        return shape

    def text(self, slide, x, y, w, h, value: Any, size=18, color=None, bold=False, align=None, wrap=True):
        from pptx.util import Pt

        frame = slide.shapes.add_textbox(self.Inches(x), self.Inches(y), self.Inches(w), self.Inches(h)).text_frame
        frame.word_wrap = wrap
        frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
        lines = value if isinstance(value, list) else [value]
        for index, line in enumerate(lines):
            paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
            run = paragraph.add_run()
            run.text = text(line)
            run.font.size = Pt(size)
            run.font.bold = bold
            run.font.name = FONT
            run.font.color.rgb = _rgb(color or self.ink)
            if align is not None:
                paragraph.alignment = align
            if index:
                paragraph.space_before = Pt(size * 0.5)
        return frame

    def image(self, slide, url, x, y, w=None, h=None) -> bool:
        stream = _image_stream(url, self.static_folder)
        if not stream:
            return False
        try:
            slide.shapes.add_picture(stream, self.Inches(x), self.Inches(y),
                                     self.Inches(w) if w else None, self.Inches(h) if h else None)
            return True
        except Exception:
            return False

    def cover_image(self, slide, url, x, y, w, h) -> bool:
        """Preenche a área toda e recorta o excesso, como ``object-fit: cover``."""
        stream = _image_stream(url, self.static_folder)
        if not stream:
            return False
        try:
            picture = slide.shapes.add_picture(stream, self.Inches(x), self.Inches(y), self.Inches(w), self.Inches(h))
            width, height = picture.image.size
            box, image = w / h, width / height
            if image > box:
                cut = (1 - box / image) / 2
                picture.crop_left = picture.crop_right = cut
            elif image < box:
                cut = (1 - image / box) / 2
                picture.crop_top = picture.crop_bottom = cut
            return True
        except Exception:
            return False

    def kicker(self, slide, data: dict, dark=False):
        label = f"{int(data.get('index') or 0):02d}  {text(data.get('kicker'))}".strip()
        self.text(slide, PAD, 0.55, 9, 0.4, label.upper(), size=12, bold=True, color=self.accent if dark else self.brand)

    def title(self, slide, value, y=1.0, size=34, color=None, w=W_IN - 2 * PAD):
        self.text(slide, PAD, y, w, 1.3, value, size=size, bold=True, color=color)

    def notes(self, slide, lines: list[str]):
        if lines:
            slide.notes_slide.notes_text_frame.text = "\n\n".join(text(line) for line in lines if text(line))

    def footer(self, slide, view: dict, data: dict):
        self.text(slide, PAD, H_IN - 0.5, 8, 0.3, text(view.get("client") or view.get("title")), size=10, color=self.muted)
        self.text(slide, W_IN - PAD - 1, H_IN - 0.5, 1, 0.3, str(data.get("index") or ""), size=10, color=self.muted)


def _cover(deck: _Deck, data: dict, view: dict, template: dict):
    slide = deck.slide(deck.ink)
    if not deck.image(slide, template.get("logo"), PAD, 0.55, h=0.45):
        deck.text(slide, PAD, 0.55, 5, 0.4, text(template.get("name")), size=14, bold=True, color="#FFFFFF")
    deck.text(slide, PAD, 1.6, 7.6, 0.4, text(data.get("label")).upper(), size=12, bold=True, color=deck.accent)
    deck.text(slide, PAD, 2.1, 7.6, 2.4, data.get("title"), size=40, bold=True, color="#FFFFFF")
    who = " · ".join(item for item in (text(data.get("client")), text(data.get("agency"))) if item)
    if who:
        deck.text(slide, PAD, 4.6, 7.6, 0.4, who, size=16, color="#FFFFFF")
    facts = as_list(data.get("facts"))
    for index, fact in enumerate(facts[:4]):
        fact = as_dict(fact)
        x = PAD + index * 1.95
        deck.text(slide, x, 5.6, 1.8, 0.3, text(fact.get("label")).upper(), size=10, bold=True, color=deck.accent)
        deck.text(slide, x, 5.9, 1.8, 0.9, fact.get("value"), size=13, color="#FFFFFF")
    if not deck.cover_image(slide, data.get("image"), 8.75, 0, W_IN - 8.75, H_IN):
        deck.box(slide, 8.75, 0, W_IN - 8.75, H_IN, deck.brand)
    deck.image(slide, data.get("logo"), 9.3, H_IN - 1.3, h=0.7)
    return slide


def _thesis(deck, data, view):
    slide = deck.slide()
    deck.kicker(slide, data)
    deck.box(slide, PAD, 1.25, 0.08, 2.6, deck.accent)
    deck.text(slide, PAD + 0.35, 1.2, 10.8, 3.2, data.get("statement"), size=32, bold=True)
    if text(data.get("support")):
        deck.text(slide, PAD + 0.35, 4.7, 10.5, 1.8, data.get("support"), size=18, color=deck.muted)
    return slide


def _body_size(body, width: float) -> float:
    """Corpo menor quando o texto é longo para a coluna; o card não pode estourar."""
    chars = sum(len(text(line)) for line in (body if isinstance(body, list) else [body]))
    density = chars / max(width, 1)
    return 15 if density < 70 else 13.5 if density < 110 else 12


def _columns(deck, slide, blocks: list[dict], y=1.3, height=4.8):
    count = max(1, len(blocks))
    gap = 0.35
    width = (W_IN - 2 * PAD - gap * (count - 1)) / count
    for index, block in enumerate(blocks):
        x = PAD + index * (width + gap)
        deck.box(slide, x, y, width, height, deck.soft)
        inner = x + 0.35
        if text(block.get("kicker")):
            deck.text(slide, inner, y + 0.3, width - 0.7, 0.3, text(block.get("kicker")).upper(), size=11, bold=True, color=deck.brand)
        top = y + 0.7
        if text(block.get("title")):
            deck.text(slide, inner, top, width - 0.7, 1.1, block.get("title"), size=22, bold=True)
            top += 1.2
        body = block.get("body")
        if body:
            deck.text(slide, inner, top, width - 0.7, y + height - top - 0.2, body, size=_body_size(body, width), color=deck.muted)


def _context(deck, data, view):
    slide = deck.slide()
    deck.kicker(slide, data)
    _columns(deck, slide, [as_dict(b) for b in as_list(data.get("blocks"))])
    return slide


def _audience(deck, data, view):
    slide = deck.slide()
    deck.kicker(slide, data)
    deck.text(slide, PAD, 1.1, W_IN - 2 * PAD, 1.4, data.get("audience"), size=22, bold=True)
    steps = [as_dict(s) for s in as_list(data.get("steps"))]
    if steps:
        _columns(deck, slide, steps, y=3.0, height=3.4)
    return slide


def _creative(deck, data, view):
    slide = deck.slide()
    deck.kicker(slide, data)
    has_image = deck.cover_image(slide, data.get("image"), 7.4, 1.1, 5.2, 5.2 * 9 / 16)
    width = 6.2 if has_image else W_IN - 2 * PAD
    deck.text(slide, PAD, 1.1, width, 1.6, data.get("title"), size=28, bold=True)
    deck.text(slide, PAD, 2.8, width, 1.6, data.get("body"), size=16, color=deck.muted)
    for index, item in enumerate(as_list(data.get("outputs"))[:3]):
        item = as_dict(item)
        y = 4.6 + index * 0.75
        deck.box(slide, PAD, y + 0.05, 0.08, 0.5, deck.accent)
        deck.text(slide, PAD + 0.25, y, width - 0.3, 0.3, item.get("name"), size=14, bold=True)
        deck.text(slide, PAD + 0.25, y + 0.3, width - 0.3, 0.4, item.get("body"), size=12, color=deck.muted)
    return slide


def _share_bar(deck, slide, share: list[dict], y: float):
    total = sum(max(0, int(as_dict(s).get("pct") or 0)) for s in share) or 1
    x = PAD
    width = W_IN - 2 * PAD
    for index, item in enumerate(share):
        item = as_dict(item)
        part = width * max(0, int(item.get("pct") or 0)) / total
        if part <= 0:
            continue
        color = deck.chart[index % len(deck.chart)]
        deck.box(slide, x, y, part, 0.42, color)
        label = f"{text(item.get('label'))} · {int(item.get('pct') or 0)}%"
        deck.text(slide, x, y + 0.5, max(part, 1.2), 0.3, label, size=12, bold=True)
        x += part


def _media(deck, data, view):
    slide = deck.slide()
    deck.kicker(slide, data)
    deck.title(slide, data.get("title"), size=28)
    share = [as_dict(s) for s in as_list(data.get("share"))]
    if share:
        _share_bar(deck, slide, share, 2.2)
    fronts = [as_dict(f) for f in as_list(data.get("fronts"))]
    if fronts:
        blocks = []
        for front in fronts:
            weight = f" · {front['weight']}%" if front.get("weight") else ""
            tags = " · ".join(text(t) for t in as_list(front.get("tags")) if text(t))
            blocks.append({"kicker": f"{text(front.get('label'))}{weight}", "title": "", "body": [front.get("text"), tags] if tags else front.get("text")})
        _columns(deck, slide, blocks, y=3.3, height=2.9)
    else:
        for index, channel in enumerate(as_list(data.get("channels"))[:6]):
            channel = as_dict(channel)
            y = 3.3 + index * 0.5
            deck.box(slide, PAD, y + 0.08, 0.22, 0.22, deck.chart[index % len(deck.chart)])
            deck.text(slide, PAD + 0.4, y, 3.4, 0.4, channel.get("label"), size=15, bold=True)
            deck.text(slide, PAD + 3.9, y, 6.6, 0.4, channel.get("role"), size=13, color=deck.muted)
            deck.text(slide, W_IN - PAD - 1, y, 1, 0.4, f"{channel.get('pct') or 0}%", size=15, bold=True)
    return slide


def _formats(deck, data, view):
    slide = deck.slide()
    deck.kicker(slide, data)
    blocks = []
    for group in as_list(data.get("groups"))[:3]:
        group = as_dict(group)
        names = [("◆ " if as_dict(f).get("aggregated") else "") + text(as_dict(f).get("name")) for f in as_list(group.get("formats"))]
        if group.get("more"):
            names.append(f"+{group['more']} formatos")
        blocks.append({"kicker": "Formatos", "title": group.get("channel"), "body": names})
    sites = as_list(data.get("sites"))
    if sites:
        head = f"Rede de portais · {data.get('sites_total')} sites" if data.get("sites_total") else "Rede de portais"
        body = [", ".join(text(s) for s in sites)]
        if text(data.get("segment")):
            body.append(f"Recorte: {text(data.get('segment'))}")
        blocks.append({"kicker": "Principais sites", "title": head, "body": body})
    _columns(deck, slide, blocks[:4], y=1.1, height=5.0)
    if as_list(data.get("aggregated")):
        deck.text(slide, PAD, 6.35, W_IN - 2 * PAD, 0.4,
                  "◆ Formatos com custo agregado, a orçar: " + ", ".join(text(x) for x in as_list(data.get("aggregated"))),
                  size=11, color=deck.muted)
    return slide


def _impact(deck, data, view):
    from pptx.util import Pt

    slide = deck.slide()
    deck.kicker(slide, data)
    if text(data.get("origin")):
        deck.text(slide, W_IN - PAD - 5, 0.55, 5, 0.4, data.get("origin"), size=12, color=deck.muted)
    cards = [as_dict(c) for c in as_list(data.get("cards"))[:4]]
    if cards:
        width = (W_IN - 2 * PAD - 0.3 * (len(cards) - 1)) / len(cards)
        for index, card in enumerate(cards):
            x = PAD + index * (width + 0.3)
            deck.box(slide, x, 1.1, width, 1.4, deck.soft)
            deck.text(slide, x + 0.25, 1.25, width - 0.5, 0.3, text(card.get("label")).upper(), size=10, bold=True, color=deck.brand)
            deck.text(slide, x + 0.25, 1.55, width - 0.5, 0.5, card.get("value"), size=26, bold=True)
            if text(card.get("note")):
                deck.text(slide, x + 0.25, 2.1, width - 0.5, 0.3, card.get("note"), size=10, color=deck.muted)
    rows = [as_dict(r) for r in as_list(data.get("rows"))[:5]]
    if rows:
        headers = ["Canal", "Investimento", "CPM", "Impressões", "Cliques"]
        shape = slide.shapes.add_table(len(rows) + 1, 5, deck.Inches(PAD), deck.Inches(2.8), deck.Inches(W_IN - 2 * PAD), deck.Inches(0.4 * (len(rows) + 1)))
        table = shape.table
        widths = (3.4, 2.2, 1.6, 2.2, 2.4)
        scale = (W_IN - 2 * PAD) / sum(widths)
        for index, inches in enumerate(widths):
            table.columns[index].width = deck.Inches(inches * scale)
        values = [headers] + [[r.get("label"), r.get("invested"), r.get("cpm") or "—", r.get("impressions") or "—", r.get("clicks") or "—"] for r in rows]
        for r_index, row in enumerate(values):
            for c_index, value in enumerate(row):
                cell = table.cell(r_index, c_index)
                cell.text = text(value)
                cell.fill.solid()
                cell.fill.fore_color.rgb = _rgb(deck.brand if r_index == 0 else deck.paper)
                font = cell.text_frame.paragraphs[0].runs[0].font if cell.text_frame.paragraphs[0].runs else None
                if font:
                    font.size = Pt(12)
                    font.name = FONT
                    font.bold = r_index == 0 or c_index == 0
                    font.color.rgb = _rgb("#FFFFFF" if r_index == 0 else deck.ink)
    indicator = as_dict(data.get("indicator"))
    if indicator.get("label"):
        line = f"{text(indicator.get('label'))}: meta {text(indicator.get('target')) or '—'} · estimado {text(indicator.get('estimated')) or '—'}"
        deck.text(slide, PAD, 5.6, W_IN - 2 * PAD, 0.4, line, size=14, bold=True)
    deck.text(slide, PAD, 6.3, W_IN - 2 * PAD, 0.5, data.get("note"), size=10, color=deck.muted)
    return slide


def _defense(deck, data, view):
    slide = deck.slide()
    deck.kicker(slide, data)
    blocks = [{"kicker": "", "title": as_dict(g).get("title"), "body": ["• " + text(i) for i in as_list(as_dict(g).get("items"))]} for g in as_list(data.get("groups"))]
    _columns(deck, slide, blocks, y=1.1, height=5.2)
    return slide


def _market(deck, data, view):
    slide = deck.slide()
    deck.kicker(slide, data)
    deck.text(slide, PAD, 1.3, 4.6, 1.8, data.get("stat"), size=60, bold=True, color=deck.brand)
    deck.text(slide, 5.7, 1.3, 6.9, 1.0, data.get("title"), size=26, bold=True)
    deck.text(slide, 5.7, 2.5, 6.9, 3.0, data.get("insight"), size=17, color=deck.muted)
    if text(data.get("source")):
        deck.text(slide, 5.7, 6.0, 6.9, 0.4, "Fonte: " + text(data.get("source")), size=11, color=deck.muted)
    return slide


def _contact(deck, data, view):
    slide = deck.slide(deck.ink)
    deck.text(slide, PAD, 0.55, 9, 0.4, text(data.get("kicker")).upper(), size=12, bold=True, color=deck.accent)
    deck.text(slide, PAD, 2.2, 10, 1.4, "Vamos alinhar os próximos passos", size=40, bold=True, color="#FFFFFF")
    if text(data.get("name")):
        deck.text(slide, PAD, 4.0, 10, 0.5, text(data.get("name")), size=20, bold=True, color="#FFFFFF")
    links = [text(data.get("whatsapp")), text(view.get("document_url"))]
    deck.text(slide, PAD, 4.7, 11, 1.0, [link for link in links if link], size=13, color="#FFFFFF")
    return slide


RENDERERS = {
    "thesis": _thesis, "context": _context, "audience": _audience, "creative": _creative,
    "media": _media, "formats": _formats, "impact": _impact, "defense": _defense,
    "market": _market, "contact": _contact,
}


def build_pptx(view: dict, slides: list[dict], template: dict, static_folder: str = "") -> bytes:
    deck = _Deck(template, static_folder)
    for data in slides:
        data = as_dict(data)
        kind = text(data.get("type"))
        if kind == "cover":
            slide = _cover(deck, data, view, template)
        elif kind in RENDERERS:
            slide = RENDERERS[kind](deck, data, view)
            deck.footer(slide, view, data)
        else:
            continue
        deck.notes(slide, as_list(data.get("notes")))
    out = io.BytesIO()
    deck.prs.save(out)
    return out.getvalue()
