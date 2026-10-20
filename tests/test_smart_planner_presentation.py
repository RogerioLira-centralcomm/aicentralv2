import base64
import io

import pytest

from aicentralv2.smart_planner import presentation
from aicentralv2.smart_planner.images import ART_PUBLIC_PREFIX, is_generated_art


def _parties(agency_brand=None, client_brand=None):
    return {
        "client": {"id": 10, "name": "Animale", "brand": client_brand or {}},
        "agency": {"id": 20, "name": "Agência Rua", "logo_url": "", "brand": agency_brand or {}},
    }


def _row(choice=""):
    return {"dados_detectados": {"presentation_template": choice} if choice else {}}


def test_hex_colors_reads_workspace_palette_records_and_text():
    palette = [{"hex": "#0b6b52", "name": "Verde"}, "#F0B000", {"hex": "nope"}, "#abc", "#0B6B52"]
    assert presentation.hex_colors(palette) == ["#0B6B52", "#F0B000", "#AABBCC"]


def test_palette_tokens_keep_readable_contrast_even_for_light_brand_colors():
    tokens = presentation.palette_tokens(["#FFE14D", "#7FD1FF"])
    assert presentation.contrast(tokens["ink"], tokens["paper"]) >= 10
    assert presentation.contrast(tokens["brand"], tokens["paper"]) >= 4.5
    assert presentation.contrast(tokens["accent"], tokens["ink"]) >= 3
    assert presentation.contrast(tokens["muted"], tokens["paper"]) >= 4.5


def test_centralcomm_template_uses_the_logo_colors():
    tokens = presentation.CENTRALCOMM["tokens"]
    assert tokens["accent"] == "#F0B000"
    assert tokens["brand"] == "#0B6B52"


def test_brand_without_logo_or_palette_gives_no_template():
    assert presentation.from_brand("agencia", {"name": "Sem marca"}) is None


def test_agency_template_is_the_default_when_it_has_a_palette():
    parties = _parties(agency_brand={"name": "Agência Rua", "palette": ["#C8102E", "#111111"], "logo_url": "/static/x.png"})
    template = presentation.resolve(_row(), parties=parties)
    assert template["id"] == "agencia"
    assert template["chosen"] is False
    assert template["tokens"]["brand"].startswith("#")


def test_without_agency_brand_default_is_centralcomm_and_choice_is_respected():
    parties = _parties(client_brand={"name": "Animale", "palette": ["#000000", "#B08D57"]})
    assert presentation.resolve(_row(), parties=parties)["id"] == "centralcomm"
    chosen = presentation.resolve(_row("cliente"), parties=parties)
    assert chosen["id"] == "cliente" and chosen["chosen"] is True


def test_neutral_palette_keeps_logo_and_turns_colors_gray():
    parties = _parties(agency_brand={"name": "Agência Rua", "palette": ["#C8102E"], "logo_url": "/static/x.png"})
    template = presentation.resolve(_row(), neutral_palette=True, parties=parties)
    assert template["neutral"] is True
    assert template["logo"] == "/static/x.png"
    assert template["tokens"] == presentation.NEUTRO["tokens"]


def test_css_vars_only_accepts_hex_values():
    template = {"tokens": {"ink": "#10252D", "brand": "red;}</style><script>", "accent": "#F0B000"}, "chart": ["#0B6B52", "url(x)"]}
    css = presentation.css_vars(template)
    assert "--sx-ink:#10252D" in css and "--sx-accent:#F0B000" in css and "--sx-c1:#0B6B52" in css
    assert "script" not in css and "url" not in css and "--sx-brand" not in css


def test_apply_to_view_recolors_channels_and_uses_agency_logo_in_the_bar():
    template = presentation.from_brand("agencia", {"name": "Agência Rua", "palette": ["#C8102E", "#1D1D1B"], "logo_url": "/static/rua.png"})
    view = {"house": {"name": "Centralcomm", "logo": "/static/images/cc_logo.png"},
            "hero": {}, "media": {"channels": [{"label": "Meta", "color": "#1877F2"}, {"label": "Rede", "color": "#000"}]}}
    out = presentation.apply_to_view(view, template)
    assert out["house"]["logo"] == "/static/rua.png"
    assert [c["color"] for c in out["media"]["channels"]] == template["chart"][:2]
    assert out["hero"]["theme_accent"] == template["tokens"]["accent"]
    assert "--ink:" in out["document_template_css"]


def test_colors_from_image_finds_the_logo_colors_and_ignores_white():
    from PIL import Image

    image = Image.new("RGB", (60, 60), "#FFFFFF")
    for x in range(10, 30):
        for y in range(10, 50):
            image.putpixel((x, y), (200, 16, 46))
    for x in range(32, 50):
        for y in range(10, 50):
            image.putpixel((x, y), (0, 90, 70))
    buffer = io.BytesIO()
    image.save(buffer, format="PNG")
    data_url = "data:image/png;base64," + base64.b64encode(buffer.getvalue()).decode()
    colors = presentation.colors_from_image(data_url)
    assert colors and "#FFFFFF" not in colors
    assert any(presentation.contrast(color, "#C8102E") < 1.3 for color in colors)
    assert presentation.colors_from_image("not-an-image") == []


def test_save_party_template_requires_a_crm_party(monkeypatch):
    monkeypatch.setattr(presentation, "plan_parties", lambda row: {"client": {}, "agency": {}})
    with pytest.raises(ValueError):
        presentation.save_party_template({}, "agencia", "", ["#000000"])
    with pytest.raises(ValueError):
        presentation.save_party_template({}, "outro", "", ["#000000"])


def test_editor_panel_lists_missing_brands_to_build(monkeypatch):
    monkeypatch.setattr(presentation, "plan_parties", lambda row: _parties())
    panel = presentation.editor_panel(_row())
    assert [o["id"] for o in panel["options"]] == ["centralcomm", "neutro"]
    assert {m["alvo"] for m in panel["missing"]} == {"agencia", "cliente"}
    assert panel["selected"] == "centralcomm"


def test_new_art_goes_to_the_folder_kept_by_deploys():
    assert ART_PUBLIC_PREFIX == "/static/uploads/smart_planner/"
    assert is_generated_art("/static/uploads/smart_planner/a.png")
    assert is_generated_art("/static/images/smart_planner/generated/a.png")
    assert not is_generated_art("/static/images/smart_planner/bg-finance.png")


def test_build_pptx_creates_editable_slides_with_notes():
    from pptx import Presentation

    from aicentralv2.smart_planner.pptx_export import build_pptx

    slides = [
        {"type": "cover", "index": 1, "title": "Animale Inverno", "client": "Animale", "label": "Plano de mídia", "facts": [{"label": "Período", "value": "Jun–Jul"}], "notes": []},
        {"type": "thesis", "index": 2, "kicker": "A decisão", "statement": "Impacto em portais AB e remarketing com DCO.", "support": "", "notes": ["Texto completo da tese."]},
        {"type": "media", "index": 3, "kicker": "Estratégia de mídia", "title": "Como a verba trabalha",
         "share": [{"label": "Impacto", "pct": 70}, {"label": "Remarketing", "pct": 30}],
         "fronts": [{"label": "Impacto", "weight": 70, "text": "Interativos na rede de portais.", "tags": ["Display"]}], "channels": [], "flight": [], "notes": []},
        {"type": "impact", "index": 4, "kicker": "Estimativa de impacto", "origin": "Estimativa de referência",
         "cards": [{"label": "Impressões", "value": "12 mi", "note": ""}],
         "rows": [{"label": "Rede de portais", "invested": "R$ 100 mil", "cpm": "R$ 12", "impressions": "8 mi", "clicks": "16 mil"}],
         "indicator": {}, "note": "Estimativa de referência.", "notes": []},
        {"type": "contact", "index": 5, "kicker": "Próximos passos", "name": "Ana", "whatsapp": "", "notes": []},
    ]
    data = build_pptx({"client": "Animale", "title": "Inverno"}, slides, presentation.CENTRALCOMM)
    deck = Presentation(io.BytesIO(data))
    assert len(deck.slides) == 5
    texts = " ".join(shape.text_frame.text for slide in deck.slides for shape in slide.shapes if shape.has_text_frame)
    assert "Animale Inverno" in texts and "Impacto · 70%" in texts
    assert deck.slides[1].notes_slide.notes_text_frame.text == "Texto completo da tese."


def test_logo_only_brand_can_still_be_completed(monkeypatch):
    monkeypatch.setattr(presentation, "plan_parties", lambda row: _parties(agency_brand={"name": "Agência Rua", "logo_url": "/static/rua.png"}))
    panel = presentation.editor_panel(_row())
    assert "agencia" in [o["id"] for o in panel["options"]]
    assert "agencia" in {m["alvo"] for m in panel["missing"]}


def test_save_never_overwrites_an_analyzed_brand(monkeypatch):
    import aicentralv2.smart_planner.brand as brand_mod

    monkeypatch.setattr(presentation, "plan_parties", lambda row: _parties())
    monkeypatch.setattr(brand_mod, "load_cx_client_for_crm", lambda crm_id: {
        "id": 5, "logo_url": "https://x/logo.png", "primary_color": "#123456",
        "brand_profile": {"color_palette": [{"hex": "#123456"}]}})
    monkeypatch.setattr(presentation, "_save_logo", lambda data: pytest.fail("não deveria gravar logo"))
    with pytest.raises(ValueError):
        presentation.save_party_template({}, "cliente", "data:image/png;base64,AAAA", ["#000000"])


def test_existing_agency_logo_prefills_colors_instead_of_asking_for_upload(monkeypatch):
    monkeypatch.setattr(presentation, "plan_parties", lambda row: {
        "client": {}, "agency": {"id": 20, "name": "Binder", "logo_url": "/static/binder.png", "brand": {}}})
    monkeypatch.setattr(presentation, "colors_from_url", lambda url: ["#111111", "#C8102E"])
    panel = presentation.editor_panel(_row())
    item = next(m for m in panel["missing"] if m["alvo"] == "agencia")
    assert item["logo"] == "/static/binder.png" and item["colors"] == ["#111111", "#C8102E"]


def test_editor_flight_bars_get_height_from_amount():
    from aicentralv2.smart_planner.editor import _media_for_editor

    media = _media_for_editor({"pace": {"months": [{"label": "S1", "amount": 50}, {"label": "S2", "amount": 100}, {"label": "S3"}]}})
    assert [m["pct"] for m in media["pace"]["months"]] == [50, 100, 0]


def test_editor_lists_formats_per_channel_from_review_and_taticas():
    from aicentralv2.smart_planner.editor import _channel_formats

    rows = _channel_formats({"campanha": {"canais": ["netflix", "dv360"], "formatos": {"dv360": ["interativo-hotspots-shop-the-look"]}}}, {"channels": [{"id": "netflix", "label": "Netflix"}]})
    by_id = {row["id"]: row for row in rows}
    assert by_id["netflix"]["formats"] == []  # sem escolha: o editor mostra quantos o catálogo oferece
    assert by_id["dv360"]["label"] == "Rede de portais e sites"
    assert by_id["dv360"]["formats"][0]["type"] == "Interativo" and by_id["dv360"]["formats"][0]["aggregated"]


def test_logo_fetch_refuses_internal_hosts():
    assert not presentation._public_host("http://169.254.169.254/latest/meta-data")
    assert not presentation._public_host("http://127.0.0.1/logo.png")
    assert not presentation._public_host("http://10.0.0.5/logo.png")


def test_executive_contact_only_in_the_centralcomm_template(monkeypatch):
    from flask import Flask

    from aicentralv2.smart_planner import routes

    agency = presentation.from_brand("agencia", {"name": "Rua", "palette": ["#C8102E"]})
    monkeypatch.setattr(presentation, "plan_parties", lambda row: _parties())
    app = Flask(__name__)
    for template, keeps in ((agency, False), (dict(presentation.CENTRALCOMM), True)):
        monkeypatch.setattr(presentation, "resolve", lambda row, neutral_palette=False, parties=None, t=template: dict(t))
        with app.test_request_context("/smart-planner/p/x/proposta"):
            view = routes._apply_presentation({"executive": {"name": "Ana", "whatsapp_url": "https://wa.me/1"}, "hero": {}, "media": {}}, {"dados_detectados": {}})
        assert bool(view["executive"]) is keeps


def test_cover_uses_the_generated_creative_as_background_with_a_mockup():
    from aicentralv2.smart_planner.public_slides import build_slides

    view = {"title": "Plano", "hero": {"creative_image": "/static/c.png", "has_generated_visual": True},
            "media": {"channels": [{"id": "netflix", "label": "Netflix", "pct": 100, "group": "ctv"}]}}
    cover = build_slides(view)[0]
    assert cover["layout"] == "hero" and cover["background"] == "/static/c.png"
    assert cover["mockup"]["kind"] == "tv"
    plain = build_slides({**view, "hero": {"image": "/static/bg.png"}})[0]
    assert plain["layout"] == ""
